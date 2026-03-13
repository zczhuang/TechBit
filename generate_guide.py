from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────
sec = doc.sections[0]
sec.page_width    = Inches(8.5)
sec.page_height   = Inches(11)
sec.left_margin   = Inches(1.1)
sec.right_margin  = Inches(1.1)
sec.top_margin    = Inches(1.0)
sec.bottom_margin = Inches(1.0)

NAVY  = RGBColor(0x1B, 0x3A, 0x6B)
GOLD  = RGBColor(0xC8, 0x96, 0x3E)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
GRAY  = RGBColor(0x6B, 0x72, 0x80)
BLACK = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0xB9, 0x1C, 0x1C)

def shade_para(para, fill_hex):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)

def run(para, text, bold=False, italic=False, color=None, size=None):
    r = para.add_run(text)
    r.bold, r.italic = bold, italic
    if color: r.font.color.rgb = color
    if size:  r.font.size = Pt(size)
    return r

def cover_banner(title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    shade_para(p, "1B3A6B")
    run(p, f"\n  {title}\n", bold=True, color=WHITE, size=20)
    run(p, f"  {subtitle}\n ", color=RGBColor(0xB3, 0xC5, 0xE0), size=11)

def part_banner(text, fill, text_color=WHITE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(10)
    shade_para(p, fill)
    run(p, f"  {text}  ", bold=True, color=text_color, size=13)
    return p

def sub_banner(text, fill="E8F4F0", text_color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(8)
    shade_para(p, fill)
    run(p, f"  {text}  ", bold=True, color=text_color or GREEN, size=11)
    return p

def step_head(num, label, badge, badge_color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    run(p, f"Step {num}  ", bold=True, color=NAVY, size=12)
    run(p, f"[{badge}]  ", bold=True, color=badge_color or GOLD, size=9)
    run(p, label, bold=True, color=NAVY, size=12)
    return p

def body(text, indent=0.3):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(indent)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(5)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = BLACK
    return p

def body_rich(parts, indent=0.3):
    """parts = list of (text, bold, color) tuples"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(indent)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(5)
    for text, bold, color in parts:
        run(p, text, bold=bold, color=color or BLACK, size=10.5)
    return p

def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.3)
    p.paragraph_format.space_after  = Pt(4)
    shade_para(p, "FEF9EC")
    run(p, "  Note: ", bold=True, color=GOLD, size=10)
    run(p, text + "  ", color=RGBColor(0x78, 0x55, 0x1F), size=10)
    return p

def link_line(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(4)
    run(p, f"→ {label}: ", bold=True, color=GOLD, size=10)
    run(p, url, italic=True, color=NAVY, size=10)
    return p

def bullet(text, prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Inches(0.45)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(3)
    if prefix:
        run(p, prefix + "  ", bold=True, color=NAVY, size=10.5)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = BLACK
    return p

def divider(light=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run("─" * 74)
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor(0xE5, 0xE7, 0xEB) if light else RGBColor(0xD1, 0xD5, 0xDB)
    return p

def spacer(pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(pt)
    p.paragraph_format.space_after  = Pt(0)
    return p

def intro_para(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(10)
    for r in p.runs:
        r.font.size = Pt(10.5)
        r.font.color.rgb = GRAY
    return p

# ══════════════════════════════════════════════════════════════════
# COVER
# ══════════════════════════════════════════════════════════════════
cover_banner(
    "Getting Started with Claude Code",
    "A Complete Guide: Setup · Build a Website · Deploy on GitHub · Buy a Domain"
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run(p, "Created with Claude Code  ·  anthropic.com", italic=True, color=GRAY, size=9)
divider()

# ══════════════════════════════════════════════════════════════════
# PART 1 — SET UP CLAUDE CODE
# ══════════════════════════════════════════════════════════════════
part_banner("PART 1 — Set Up Claude Code", "1B3A6B")
intro_para(
    "Get Claude Code running on your computer in five steps. "
    "No command line or prior coding experience required."
)

step_head(1, "Subscribe to Claude Pro", "Web")
body("Open your browser and go to claude.ai. Click Sign Up to create a free Anthropic account "
     "using your email address. Once logged in, click Upgrade in the top-right corner and select "
     "the Pro plan for $20/month. This gives you access to Claude's most capable models and "
     "unlocks Claude Code for AI-powered development.")
link_line("claude.ai", "https://claude.ai")
divider()

step_head(2, "Install Claude Desktop", "Desktop")
body("Go to claude.ai/download in your browser. Click the download button for your operating "
     "system (Mac or Windows). Once downloaded, open the installer and follow the on-screen "
     "steps — just like installing any other app. Launch Claude Desktop and sign in with the "
     "same Anthropic account you just created. Claude Desktop is a full AI assistant you can "
     "use anytime without opening a browser.")
link_line("Download Claude Desktop", "https://claude.ai/download")
divider()

step_head(3, "Install Your Code Editor (IDE)", "IDE")
body("Download one of these two AI-powered code editors — both are free and beginner-friendly:")
bullet("Antigravity (by Google) — download from Google's website and install like any app.")
bullet("Cursor — go to cursor.com, click Download, and install like any app.")
body("Both have a visual, point-and-click interface. No command line or terminal is required.")
link_line("Download Cursor", "https://cursor.com")
divider()

step_head(4, "Find and Install the Claude Code Extension", "IDE")
body("Open your IDE (Antigravity or Cursor). Look at the left sidebar — click the Extensions "
     "icon (it looks like four small squares or a puzzle piece). A search panel will open. "
     "Type  Claude Code  in the search box. Click the Claude Code result from Anthropic, "
     "then click the blue Install button. The extension connects your editor directly to Claude.")
note("If you don't see the Extensions icon, go to the top menu: View → Extensions.")
divider()

step_head(5, "Log In and Activate Claude Code", "IDE")
body("After the extension installs, a Claude Code panel will appear in your IDE sidebar — "
     "click it to open it. You'll be prompted to sign in. Enter your Anthropic account "
     "credentials (the same email and password you use on claude.ai). Once signed in, "
     "Claude Code is fully active and ready to help you build.")

# ══════════════════════════════════════════════════════════════════
# PART 2 — BUILD A WEBSITE
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
part_banner("PART 2 — Build a Website with Claude Code", "2D6A4F")
intro_para(
    "You don't need to know HTML, CSS, or JavaScript. Describe what you want in plain "
    "English and Claude Code writes the code for you. Here's the complete workflow."
)

step_head(6, "Create a New Project Folder", "IDE")
body("In your IDE, go to File → Open Folder (or File → Open… on Mac). "
     "Navigate to your Desktop or Documents and create a new empty folder — "
     "name it something like my-website or campaign-site. Open that folder in the IDE. "
     "This folder is where all your website files will be saved.")
divider()

step_head(7, "Open the Claude Code Chat Panel", "IDE")
body("Click the Claude Code icon in the left sidebar to open the chat panel. "
     "You'll see a text input box — this is where you give Claude instructions. "
     "It works exactly like texting: type what you want, press Enter.")
divider()

step_head(8, "Describe Your Website", "IDE")
body("Type a description of the website you want to build. Be as specific as possible — "
     "include the purpose, sections, colors, and any content. Example prompt:")
p = doc.add_paragraph(
    '"Build me a single-page campaign website for Jane Smith running for Belmont Town '
    'Meeting, Precinct 4. Include: a full-width hero section with her name and tagline '
    '"Thoughtful Leadership for Belmont", an About section with a photo placeholder, '
    'a Priorities section with 3 cards (Education, Fiscal Responsibility, Civic Engagement), '
    'a Get Involved contact form, and a footer. Use a navy blue and warm gold color scheme '
    'with Playfair Display for headings."'
)
p.paragraph_format.left_indent = Inches(0.4)
p.paragraph_format.space_after = Pt(6)
for r in p.runs:
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY
    r.italic = True
body("Claude Code will create all the files — index.html, style.css, main.js — automatically "
     "in your project folder.")
divider()

step_head(9, "Preview Your Site in the Browser", "IDE")
body("In the IDE file explorer (left panel), right-click on index.html and choose "
     "Open with Live Server  or  Reveal in Finder/Explorer then double-click the file. "
     "Your website will open in your default browser. Live Server will auto-refresh "
     "every time Claude makes a change.")
note("If 'Open with Live Server' isn't available, install the Live Server extension "
     "the same way you installed Claude Code.")
divider()

step_head(10, "Refine with Follow-Up Instructions", "IDE")
body("Don't like something? Just type a follow-up in the Claude Code panel. "
     "No need to touch the code yourself. Examples:")
bullet('"Change the hero background to a deeper navy gradient."')
bullet('"Add a fifth step to the setup section about form submissions."')
bullet('"Make all the buttons gold with rounded corners."')
bullet('"Add a photo of Jane to the About section — the file is called jane.jpg."')
bullet('"Make the site look good on a phone screen."')
body("Keep refining until the site looks and works exactly the way you want.")

# ══════════════════════════════════════════════════════════════════
# PART 3 — DEPLOY ON GITHUB PAGES
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
part_banner("PART 3 — Deploy on GitHub Pages (Free)", "1B3A6B")
intro_para(
    "GitHub Pages hosts your website for free. Your site gets a live URL like "
    "yourusername.github.io/my-site — shareable immediately, no server or hosting fees."
)

sub_banner("Section A: Create Your GitHub Account & Repository", "EBF0F8", NAVY)

step_head("A1", "Create a Free GitHub Account", "Web")
body("Open your browser and go to github.com. Click Sign up in the top-right corner. "
     "Enter your email address, create a password, and choose a username "
     "(this will appear in your site URL, e.g. github.com/janesmithbelmont). "
     "Follow the verification steps and complete the account setup.")
link_line("github.com", "https://github.com")
divider()

step_head("A2", "Create a New Repository", "Web")
body("Once logged in, click the green New button on the left, or click the + icon "
     "in the top-right and choose New repository. Fill in the form:")
bullet("Repository name: type something like  my-campaign-site  (no spaces — use hyphens).")
bullet("Visibility: select Public (required for free GitHub Pages hosting).")
bullet("Leave all other checkboxes unchecked.")
body("Click the green Create repository button at the bottom.")
divider()

step_head("A3", "Upload Your Website Files", "Web")
body("On your new empty repository page, click Add file → Upload files. "
     "A file upload area will appear. Open your project folder on your computer "
     "and drag all your files into the browser upload area:")
bullet("index.html")
bullet("style.css")
bullet("main.js")
bullet("Any images (e.g. photo.jpg, logo.png)")
body("Scroll down to the Commit changes section. In the first text box, type a short "
     "description like  Initial website upload. Then click the green Commit changes button. "
     "Wait for the upload to complete — you'll see your files listed in the repository.")
divider()

sub_banner("Section B: Enable GitHub Pages", "EBF0F8", NAVY)

step_head("A4", "Open Repository Settings", "Web")
body("In your repository (the page showing your uploaded files), click the Settings tab "
     "at the top of the page. It has a gear icon. This opens the repository settings.")
divider()

step_head("A5", "Turn On GitHub Pages", "Web")
body("In the Settings left sidebar, scroll down and click Pages. "
     "You'll see a section called Build and deployment. Under Source, click the dropdown "
     "that says None and change it to Deploy from a branch. "
     "A second dropdown will appear — make sure it says main. Click Save.")
note("If you don't see 'main' in the branch dropdown, your files may be on a branch "
     "called 'master' instead — select that.")
divider()

step_head("A6", "Get Your Live URL", "Web")
body("After saving, wait about 2 minutes. Refresh the Settings → Pages page. "
     "You'll see a green banner that says:")
p = doc.add_paragraph(
    "Your site is live at https://yourusername.github.io/my-campaign-site"
)
p.paragraph_format.left_indent = Inches(0.4)
p.paragraph_format.space_after = Pt(6)
for r in p.runs:
    r.font.size = Pt(10.5)
    r.font.color.rgb = GREEN
    r.bold = True
body("Click that link to visit your live website. Share it on social media, in emails, "
     "or on printed flyers — your site is now publicly accessible to anyone.")
divider()

step_head("A7", "Updating Your Site Later", "Web")
body("Whenever Claude Code makes changes to your files, you'll need to re-upload them to "
     "GitHub to update the live site. Go to your repository, click Add file → Upload files "
     "again, drag the updated files in (they'll replace the old ones), and commit. "
     "GitHub Pages will rebuild automatically within 1–2 minutes.")

# ══════════════════════════════════════════════════════════════════
# PART 4 — BUY A CUSTOM DOMAIN
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
part_banner("PART 4 — Buy a Custom Domain Name", "C8963E", NAVY)
intro_para(
    "A custom domain like yourname.com or yourcampaign.com makes your site look "
    "professional and is easy to share. Domains typically cost $10–15/year. "
    "Two great options: Namecheap and Cloudflare Registrar."
)

# ── Namecheap ─────────────────────────────────────────────────────
sub_banner("Option 1: Namecheap (Beginner-Friendly)", "FEF9EC", RGBColor(0x78, 0x55, 0x1F))

step_head("N1", "Search for Your Domain", "Web")
body("Go to namecheap.com. In the large search bar on the homepage, type the domain name "
     "you want — for example  meenalbagla.com  or  janesmith4belmont.com. "
     "Click Search. Namecheap will show you whether the name is available and its price. "
     "A .com domain typically costs $10–14 for the first year.")
link_line("namecheap.com", "https://namecheap.com")
divider()

step_head("N2", "Add to Cart and Check Out", "Web")
body("Click Add to Cart next to the domain you want. Click the cart icon in the top-right "
     "and then View Cart. Uncheck any add-ons you don't need (WhoisGuard privacy is free "
     "and worth keeping). Click Confirm Order. Create a Namecheap account or log in, "
     "then enter your payment information and complete the purchase.")
divider()

step_head("N3", "Find Your DNS Settings", "Web")
body("After purchase, log in to namecheap.com and go to Account → Dashboard. "
     "Click Domain List in the left sidebar. Find your domain and click Manage. "
     "Click the Advanced DNS tab. This is where you'll add the records that point "
     "your domain to your GitHub Pages site.")
divider()

step_head("N4", "Add GitHub Pages DNS Records", "Web")
body("In Advanced DNS, delete any existing A records that are already there, "
     "then add the following records. Click Add New Record for each one:")
body("Add four A Records (one for each GitHub IP address):", indent=0.2)
bullet("Type: A Record  |  Host: @  |  Value: 185.199.108.153  |  TTL: Automatic")
bullet("Type: A Record  |  Host: @  |  Value: 185.199.109.153  |  TTL: Automatic")
bullet("Type: A Record  |  Host: @  |  Value: 185.199.110.153  |  TTL: Automatic")
bullet("Type: A Record  |  Host: @  |  Value: 185.199.111.153  |  TTL: Automatic")
body("Add one CNAME Record:", indent=0.2)
bullet("Type: CNAME Record  |  Host: www  |  Value: yourusername.github.io  |  TTL: Automatic")
body("Replace  yourusername  with your actual GitHub username. Click the checkmark to save each record.")
note("DNS changes can take 15 minutes to 24 hours to fully take effect across the internet.")
divider()

step_head("N5", "Connect Domain in GitHub Pages", "Web")
body("Go back to your GitHub repository. Click Settings → Pages. "
     "In the Custom domain field, type your domain name exactly as purchased "
     "(e.g. meenalbagla.com) and click Save. Check the box for Enforce HTTPS "
     "if it appears. GitHub will verify your domain — this may take a few minutes. "
     "Once verified, your site will be accessible at your custom domain.")

spacer(16)

# ── Cloudflare ────────────────────────────────────────────────────
sub_banner("Option 2: Cloudflare Registrar (Best Price + Security)", "E8F0FE", NAVY)

step_head("C1", "Create a Cloudflare Account", "Web")
body("Go to cloudflare.com and click Sign Up. Enter your email and create a password. "
     "Cloudflare is both a domain registrar and a security/performance layer — "
     "it gives you DNS management, free HTTPS, and DDoS protection automatically.")
link_line("cloudflare.com", "https://cloudflare.com")
divider()

step_head("C2", "Register Your Domain", "Web")
body("In your Cloudflare dashboard, click the Domain Registration tab in the left sidebar, "
     "then click Register Domains. Search for your desired domain name. "
     "Cloudflare sells domains at wholesale cost (no markup) — a .com is typically $9.77/year. "
     "Select your domain and complete the purchase with a credit or debit card.")
divider()

step_head("C3", "Add DNS Records in Cloudflare", "Web")
body("After purchasing, your domain will appear in your Cloudflare dashboard. "
     "Click on your domain name, then click DNS in the left sidebar. "
     "Click Add record for each of the following:")
body("Add four A Records:", indent=0.2)
bullet("Type: A  |  Name: @  |  IPv4: 185.199.108.153  |  Proxy: DNS only (grey cloud)")
bullet("Type: A  |  Name: @  |  IPv4: 185.199.109.153  |  Proxy: DNS only (grey cloud)")
bullet("Type: A  |  Name: @  |  IPv4: 185.199.110.153  |  Proxy: DNS only (grey cloud)")
bullet("Type: A  |  Name: @  |  IPv4: 185.199.111.153  |  Proxy: DNS only (grey cloud)")
body("Add one CNAME Record:", indent=0.2)
bullet("Type: CNAME  |  Name: www  |  Target: yourusername.github.io  |  Proxy: DNS only")
note("Set each record to 'DNS only' (grey cloud icon), NOT 'Proxied' (orange cloud). "
     "Proxied mode can interfere with GitHub Pages SSL certificate verification.")
divider()

step_head("C4", "Connect Domain in GitHub Pages", "Web")
body("Go to your GitHub repository → Settings → Pages. "
     "In the Custom domain field, type your domain (e.g. meenalbagla.com) and click Save. "
     "Enable Enforce HTTPS once GitHub verifies your domain. "
     "Your site will be live at your custom domain within minutes to a few hours.")

# ══════════════════════════════════════════════════════════════════
# PART 5 — EXTRAS
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
part_banner("PART 5 — Optional Extras", "2D6A4F")

sub_banner("Receive Contact Form Submissions by Email", "E8F4F0", GREEN)

step_head("F1", "Sign Up for Formspree", "Web")
body("Go to formspree.io and create a free account. Formspree receives form submissions "
     "from your website and forwards them to your email inbox — no server needed.")
link_line("formspree.io", "https://formspree.io")
divider()

step_head("F2", "Create a New Form", "Web")
body("In your Formspree dashboard, click New Form. Give it a name (e.g. Campaign Contact Form) "
     "and enter your email address. Formspree will generate a unique form endpoint URL "
     "that looks like:  https://formspree.io/f/abcdefgh")
divider()

step_head("F3", "Connect the Form to Your Website", "IDE")
body("Open your Claude Code panel and type:")
p = doc.add_paragraph(
    '"Update the contact form\'s action attribute to https://formspree.io/f/YOUR_ID"'
)
p.paragraph_format.left_indent = Inches(0.4)
for r in p.runs:
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY
    r.italic = True
body("Replace YOUR_ID with the actual ID from Formspree. Claude Code will update the "
     "HTML for you. Re-upload the updated index.html to GitHub as described in Part 3.")
divider()

step_head("F4", "Test It", "Web")
body("Visit your live site, fill in the contact form, and click Submit. "
     "Check your email — you should receive the submission within a minute. "
     "All future form submissions will arrive directly in your inbox.")

spacer(12)
sub_banner("Connect Google Analytics (Track Visitors)", "E8F0FE", NAVY)

step_head("G1", "Create a Google Analytics Account", "Web")
body("Go to analytics.google.com and sign in with your Google account. "
     "Click Start measuring and follow the setup wizard. When asked for a website URL, "
     "enter your GitHub Pages URL or custom domain. Google will give you a "
     "Measurement ID that looks like  G-XXXXXXXXXX.")
link_line("analytics.google.com", "https://analytics.google.com")
divider()

step_head("G2", "Add the Tracking Code", "IDE")
body("In Claude Code, type:")
p = doc.add_paragraph(
    '"Add a Google Analytics tracking script to the site. My Measurement ID is G-XXXXXXXXXX."'
)
p.paragraph_format.left_indent = Inches(0.4)
for r in p.runs:
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY
    r.italic = True
body("Claude Code will add the correct script tag to your HTML automatically. "
     "Re-upload the file to GitHub. Within 24 hours you'll start seeing visitor data "
     "in your Google Analytics dashboard.")

# ══════════════════════════════════════════════════════════════════
# QUICK REFERENCE
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
part_banner("Quick Reference — All Tools & Links", "1B3A6B")

spacer(4)
refs = [
    ("Claude Pro ($20/month)",             "claude.ai",                     "https://claude.ai"),
    ("Claude Desktop (download)",          "claude.ai/download",            "https://claude.ai/download"),
    ("Cursor IDE (free download)",         "cursor.com",                    "https://cursor.com"),
    ("GitHub (free hosting)",              "github.com",                    "https://github.com"),
    ("Namecheap (domains ~$11/yr)",        "namecheap.com",                 "https://namecheap.com"),
    ("Cloudflare Registrar (~$10/yr)",     "cloudflare.com",                "https://cloudflare.com"),
    ("Formspree (form emails, free tier)", "formspree.io",                  "https://formspree.io"),
    ("Google Analytics (free)",            "analytics.google.com",          "https://analytics.google.com"),
    ("GitHub Pages IP: 185.199.108.153",   "Add as A record in DNS",        ""),
    ("GitHub Pages IP: 185.199.109.153",   "Add as A record in DNS",        ""),
    ("GitHub Pages IP: 185.199.110.153",   "Add as A record in DNS",        ""),
    ("GitHub Pages IP: 185.199.111.153",   "Add as A record in DNS",        ""),
]

for label, display, url in refs:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    run(p, f"{label:<40}", bold=True, color=NAVY, size=10.5)
    run(p, f"  {display}", italic=True, color=GRAY, size=10)

divider()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Built with Claude Code · Anthropic  |  anthropic.com",
    italic=True, color=GRAY, size=9)

doc.save("/home/user/TechBit/Claude_Code_Guide.docx")
print("Saved: Claude_Code_Guide.docx")
